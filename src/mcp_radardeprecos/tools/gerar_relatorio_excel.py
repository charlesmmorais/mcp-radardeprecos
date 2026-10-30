"""Tool MCP: gera uma planilha Excel real (.xlsx) a partir de uma pesquisa já formalizada.

Consome o dicionário devolvido por ``gerar_memoria_pesquisa`` e o devolvido
por ``calcular_preco_referencia`` e produz um ``.xlsx`` com 4 abas.

Suporta múltiplos **modelos** de relatório (parâmetro `modelo`, ver
``domain.dominios.MODELOS_RELATORIO_PRECOS``): ``"in65"`` (padrão, único
implementado — mapa de pesquisa conforme os elementos do Art. 3º da IN
SEGES/ME nº 65/2021), ``"tic_licitacao"`` e ``"tic_inexigibilidade"`` (modelos
específicos para contratação de TIC via licitação e via
inexigibilidade/dispensa — preparados como valores válidos, mas ainda não
implementados; selecioná-los levanta ``NotImplementedError``).

Assim como ``calcular_preco_referencia``, ``classificar_curva_abc`` e
``gerar_memoria_pesquisa``, esta é uma tool **quase pura**: não consulta
nenhuma API — sua única operação com efeito colateral é gravar o arquivo no
caminho informado.

A coluna "URL / Evidência de Acesso" da aba Fontes e
Evidências mostra o ``url`` da fonte (ex.: link da página do
PNCP/Compras.gov.br do contrato/ata/edital) ou, na ausência dele, o texto
livre de ``evidencia_acesso`` (para fontes sem link, como cotação direta);
se nenhum dos dois estiver preenchido, a célula é destacada em vermelho com
um marcador de pendência — nunca fica em branco silenciosamente. Pedido
explícito do usuário: todo relatório precisa trazer o link do
PNCP/Compras.gov.br do contrato/ata ou, quando não houver, evidência clara
de como acessar a fonte.
"""
from __future__ import annotations

from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..domain.dominios import MODELOS_RELATORIO_PRECOS, MODELOS_RELATORIO_PRECOS_IMPLEMENTADOS
from ..logging_utils import get_logger
from ..utils.arquivo_saida import gravar_documento

logger = get_logger(__name__)

_CAMPOS_OBRIGATORIOS_MEMORIA = (
    "objeto",
    "agente_responsavel",
    "data_pesquisa",
    "metodologia",
    "fontes_consultadas",
    "quantidade_fontes",
)
_CAMPOS_OBRIGATORIOS_RESULTADO_CALCULO = (
    "metodo_utilizado",
    "preco_referencia",
    "estatisticas",
    "amostra_original",
    "amostra_saneada",
    "valores_excluidos",
    "iteracoes_saneamento",
)

_FILL_CABECALHO = PatternFill("solid", fgColor="1F4E78")
_FONTE_CABECALHO = Font(color="FFFFFF", bold=True)
_FILL_INCLUIDO = PatternFill("solid", fgColor="C6EFCE")
_FILL_EXCLUIDO = PatternFill("solid", fgColor="FFC7CE")
_FONTE_TITULO = Font(bold=True, size=13)
_FONTE_ROTULO = Font(bold=True)

_MARCADOR_SEM_EVIDENCIA = "[PENDÊNCIA: informar URL (PNCP/Compras.gov.br) ou evidência de acesso]"


def _url_ou_evidencia(fonte: dict[str, Any]) -> str:
    """URL da fonte (preferencial) ou o texto livre de `evidencia_acesso`;
    se nenhum dos dois estiver preenchido, devolve um marcador de pendência
    — a célula é destacada em vermelho pelo chamador (nunca fica em branco
    silenciosamente). Art. 3º, III, da IN SEGES/ME nº 65/2021 exige
    caracterizar as fontes consultadas."""
    return fonte.get("url") or fonte.get("evidencia_acesso") or _MARCADOR_SEM_EVIDENCIA


def _validar_memoria(memoria: dict[str, Any]) -> None:
    faltantes = [c for c in _CAMPOS_OBRIGATORIOS_MEMORIA if c not in memoria]
    if faltantes:
        raise ValueError(
            f"memoria não parece vir de gerar_memoria_pesquisa (faltam: {faltantes})"
        )


def _validar_resultado_calculo(resultado_calculo: dict[str, Any]) -> None:
    faltantes = [c for c in _CAMPOS_OBRIGATORIOS_RESULTADO_CALCULO if c not in resultado_calculo]
    if faltantes:
        raise ValueError(
            f"resultado_calculo não parece vir de calcular_preco_referencia (faltam: {faltantes})"
        )


def _validar_modelo(modelo: str) -> None:
    if modelo not in MODELOS_RELATORIO_PRECOS:
        raise ValueError(f"modelo deve ser um de {MODELOS_RELATORIO_PRECOS}")
    if modelo not in MODELOS_RELATORIO_PRECOS_IMPLEMENTADOS:
        raise NotImplementedError(
            f"modelo={modelo!r} está preparado (é um valor válido de "
            "domain.dominios.MODELOS_RELATORIO_PRECOS) mas ainda não foi "
            "implementado — só 'in65' gera relatório hoje. Ver README/CHANGELOG."
        )


def _cabecalho(ws: Worksheet, linha: int, colunas: list[str]) -> None:
    for col, titulo in enumerate(colunas, start=1):
        celula = ws.cell(row=linha, column=col, value=titulo)
        celula.font = _FONTE_CABECALHO
        celula.fill = _FILL_CABECALHO
        celula.alignment = Alignment(wrap_text=True, vertical="center")


def _autoajustar_largura(ws: Worksheet, larguras: list[int]) -> None:
    for col, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(col)].width = largura


def _montar_aba_mapa_pesquisa(wb: Workbook, memoria: dict[str, Any], resultado_calculo: dict[str, Any]) -> None:
    ws = wb.active
    ws.title = "Mapa de Pesquisa"

    ws["A1"] = "Mapa de Pesquisa de Preços"
    ws["A1"].font = _FONTE_TITULO
    ws.merge_cells("A1:B1")

    metodologia = memoria.get("metodologia") or {}
    estatisticas = resultado_calculo.get("estatisticas") or {}

    linhas: list[tuple[str, Any]] = [
        ("Objeto", memoria.get("objeto")),
        ("Agente responsável", memoria.get("agente_responsavel")),
        ("Data da pesquisa", memoria.get("data_pesquisa")),
        ("Método de cálculo utilizado", resultado_calculo.get("metodo_utilizado") or metodologia.get("metodo_utilizado")),
        ("Preço de referência (R$)", resultado_calculo.get("preco_referencia") or metodologia.get("preco_referencia")),
        ("Quantidade de fontes", memoria.get("quantidade_fontes")),
        ("Buscas frustradas registradas", len(memoria.get("buscas_frustradas") or [])),
        ("", ""),
        ("Média da amostra final (R$)", estatisticas.get("media")),
        ("Mediana da amostra final (R$)", estatisticas.get("mediana")),
        ("Mínimo da amostra final (R$)", estatisticas.get("minimo")),
        ("Máximo da amostra final (R$)", estatisticas.get("maximo")),
        ("Desvio-padrão da amostra final", estatisticas.get("desvio_padrao")),
        ("Coeficiente de Variação (%)", estatisticas.get("coeficiente_variacao_percentual")),
        ("Iterações de saneamento aplicadas", resultado_calculo.get("iteracoes_saneamento")),
        ("Valores excluídos no saneamento", len(resultado_calculo.get("valores_excluidos") or [])),
    ]

    linha_atual = 3
    for rotulo, valor in linhas:
        ws.cell(row=linha_atual, column=1, value=rotulo).font = _FONTE_ROTULO
        ws.cell(row=linha_atual, column=2, value=valor)
        linha_atual += 1

    avisos = list(memoria.get("avisos") or [])
    if avisos:
        linha_atual += 1
        ws.cell(row=linha_atual, column=1, value="Avisos").font = _FONTE_ROTULO
        linha_atual += 1
        for aviso in avisos:
            ws.cell(row=linha_atual, column=1, value=f"- {aviso}")
            ws.merge_cells(start_row=linha_atual, start_column=1, end_row=linha_atual, end_column=2)
            linha_atual += 1

    _autoajustar_largura(ws, [32, 60])


def _montar_aba_fontes(wb: Workbook, memoria: dict[str, Any]) -> None:
    ws = wb.create_sheet("Fontes e Evidências")
    colunas = [
        "#", "Informante", "CNPJ", "Meio de consulta", "Valor unitário (R$)",
        "Quantidade", "Data de referência", "Especificação", "Condições",
        "URL / Evidência de Acesso",
    ]
    _cabecalho(ws, 1, colunas)

    fontes = memoria.get("fontes_consultadas") or []
    linha = 2
    for i, fonte in enumerate(fontes, start=1):
        sem_evidencia = not fonte.get("url") and not fonte.get("evidencia_acesso")
        valores = [
            i,
            fonte.get("informante"),
            fonte.get("cnpj"),
            fonte.get("meio_consulta"),
            fonte.get("valor_unitario"),
            fonte.get("quantidade"),
            fonte.get("data_referencia"),
            fonte.get("especificacao"),
            fonte.get("condicoes"),
            _url_ou_evidencia(fonte),
        ]
        for col, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha, column=col, value=valor)
            if sem_evidencia and col == len(valores):
                celula.fill = _FILL_EXCLUIDO
        linha += 1

    buscas_frustradas = memoria.get("buscas_frustradas") or []
    if buscas_frustradas:
        linha += 1
        ws.cell(row=linha, column=1, value="Buscas sem resultado").font = _FONTE_ROTULO
        linha += 1
        _cabecalho(ws, linha, ["#", "Meio de consulta", "Termo de busca", "Motivo"])
        linha += 1
        for i, busca in enumerate(buscas_frustradas, start=1):
            ws.cell(row=linha, column=1, value=i)
            ws.cell(row=linha, column=2, value=busca.get("meio_consulta"))
            ws.cell(row=linha, column=3, value=busca.get("termo_busca"))
            ws.cell(row=linha, column=4, value=busca.get("motivo"))
            linha += 1

    _autoajustar_largura(ws, [4, 26, 18, 24, 16, 12, 16, 32, 24, 40])


def _montar_aba_analise_estatistica(wb: Workbook, resultado_calculo: dict[str, Any]) -> None:
    ws = wb.create_sheet("Análise Estatística")

    ws["A1"] = "Amostra original"
    ws["A1"].font = _FONTE_ROTULO
    amostra_original = (resultado_calculo.get("amostra_original") or {}).get("valores") or []
    valores_excluidos = resultado_calculo.get("valores_excluidos") or []

    _cabecalho(ws, 2, ["#", "Valor (R$)", "Situação"])
    linha = 3
    excluidos_restantes = list(valores_excluidos)
    for i, valor in enumerate(amostra_original, start=1):
        situacao = "Considerado na média"
        # Usamos a lista de excluídos (removendo cada um ao encontrá-lo, para
        # tratar empates de valor corretamente) para decidir a situação/cor
        # da célula — não um "in" contra o set da amostra saneada, que
        # colidiria em empates.
        if valor in excluidos_restantes:
            situacao = "Excluído no saneamento"
            excluidos_restantes.remove(valor)
        ws.cell(row=linha, column=1, value=i)
        celula_valor = ws.cell(row=linha, column=2, value=valor)
        ws.cell(row=linha, column=3, value=situacao)
        celula_valor.fill = _FILL_EXCLUIDO if situacao == "Excluído no saneamento" else _FILL_INCLUIDO
        linha += 1

    linha += 1
    ws.cell(row=linha, column=1, value="Resumo do saneamento").font = _FONTE_ROTULO
    linha += 1
    estatisticas = resultado_calculo.get("estatisticas") or {}
    resumo = [
        ("Método utilizado", resultado_calculo.get("metodo_utilizado")),
        ("Iterações de saneamento", resultado_calculo.get("iteracoes_saneamento")),
        ("Quantidade original", len(amostra_original)),
        ("Quantidade após saneamento", len(resultado_calculo.get("amostra_saneada") or [])),
        ("Quantidade excluída", len(valores_excluidos)),
        ("Preço de referência (R$)", resultado_calculo.get("preco_referencia")),
        ("Média final (R$)", estatisticas.get("media")),
        ("Mediana final (R$)", estatisticas.get("mediana")),
        ("Desvio-padrão final", estatisticas.get("desvio_padrao")),
        ("CV final (%)", estatisticas.get("coeficiente_variacao_percentual")),
    ]
    for rotulo, valor in resumo:
        ws.cell(row=linha, column=1, value=rotulo).font = _FONTE_ROTULO
        ws.cell(row=linha, column=2, value=valor)
        linha += 1

    _autoajustar_largura(ws, [28, 20, 24])


def _montar_aba_avisos(wb: Workbook, memoria: dict[str, Any], resultado_calculo: dict[str, Any]) -> None:
    ws = wb.create_sheet("Avisos e Pendências")
    _cabecalho(ws, 1, ["Origem", "Mensagem"])

    linha = 2
    for aviso in memoria.get("avisos") or []:
        ws.cell(row=linha, column=1, value="Memória de pesquisa")
        ws.cell(row=linha, column=2, value=aviso)
        linha += 1
    for aviso in resultado_calculo.get("avisos") or []:
        ws.cell(row=linha, column=1, value="Cálculo do preço de referência")
        ws.cell(row=linha, column=2, value=aviso)
        linha += 1

    pendencias_padrao = [
        "Confirmar data e hora exatas de acesso a cada portal consultado.",
        "Anexar comprovantes/capturas de tela das consultas realizadas.",
        "Coletar assinatura dos responsáveis pela pesquisa antes de anexar ao processo.",
    ]
    for pendencia in pendencias_padrao:
        ws.cell(row=linha, column=1, value="Pendência de preenchimento humano")
        ws.cell(row=linha, column=2, value=pendencia)
        linha += 1

    _autoajustar_largura(ws, [28, 80])


async def gerar_relatorio_excel(
    memoria: dict[str, Any],
    resultado_calculo: dict[str, Any],
    caminho_saida: str | None = None,
    modelo: str = "in65",
) -> dict[str, Any]:
    """Gera uma planilha Excel (.xlsx) real com o mapa completo de uma pesquisa de preços.

    Complementa ``gerar_memoria_pesquisa`` (que só produz JSON/Markdown em
    texto): esta tool grava um arquivo binário de verdade, pronto para
    anexar ao processo administrativo. Para ``modelo="in65"`` (padrão,
    único implementado), a planilha tem 4 abas:

    1. **Mapa de Pesquisa** — resumo executivo (objeto, agente responsável,
       método aplicado, preço de referência, estatísticas da amostra final).
    2. **Fontes e Evidências** — cada fonte consultada com seus dados de
       rastreabilidade (incluindo URL/evidência de acesso — destacada em
       vermelho quando ausente), e as buscas que não retornaram resultado.
    3. **Análise Estatística** — a amostra original valor a valor, com cor
       indicando quais foram excluídos pelo saneamento, e o resumo
       numérico do cálculo (iterações, CV, etc.).
    4. **Avisos e Pendências** — todos os avisos gerados pelas tools de
       cálculo/formalização, mais os itens que sempre exigem preenchimento
       humano (datas/horas de acesso, comprovantes, assinaturas).

    ``modelo="tic_licitacao"`` e ``modelo="tic_inexigibilidade"`` (modelos
    específicos para contratação de TIC) estão **preparados como valores
    válidos mas ainda não implementados** — selecioná-los levanta
    ``NotImplementedError``.

    Esta tool é **quase pura**: não consulta nenhuma API — sua única
    operação com efeito colateral é gravar/serializar a planilha (ver
    ``caminho_saida`` abaixo).

    Args:
        memoria: O dicionário devolvido por ``gerar_memoria_pesquisa``
            (com ou sem ``formato="markdown"`` — o campo
            ``documento_markdown``, se presente, é ignorado aqui).
        resultado_calculo: O dicionário devolvido por
            ``calcular_preco_referencia`` para a MESMA amostra usada na
            memória — traz o detalhe de amostra original/saneada/excluída
            que ``memoria`` sozinha não preserva.
        caminho_saida: Caminho absoluto ou relativo do arquivo ``.xlsx`` a
            gravar no disco de onde o servidor roda (diretórios
            intermediários são criados automaticamente). Use este
            parâmetro quando o servidor roda em ``stdio`` (mesma máquina do
            usuário, ex.: Claude Desktop) ou via túnel a partir da própria
            máquina (ex.: ChatGPT) — nesses casos o disco do servidor É o
            disco do usuário. **Omita este parâmetro quando o servidor
            estiver hospedado remotamente** (ex.: custom connector
            compartilhado por vários usuários): a planilha é então
            serializada inteiramente em memória, sem tocar o disco do
            servidor, e devolvida em base64 (ver Returns) com um nome de
            arquivo único — cabe ao cliente decodificar e salvar do lado
            do usuário.
        modelo: Modelo de relatório — "in65" (padrão, implementado),
            "tic_licitacao" ou "tic_inexigibilidade" (preparados, ainda não
            implementados — levantam ``NotImplementedError``).

    Returns:
        Dicionário com ``modelo``, ``abas`` (nomes das 4 abas, na ordem em
        que aparecem), ``tamanho_bytes`` e mais: se ``caminho_saida`` foi
        informado, ``caminho_saida`` (caminho absoluto do arquivo
        gravado); se foi omitido, ``nome_arquivo`` (nome sugerido, com um
        UUID) e ``conteudo_base64`` (conteúdo do ``.xlsx`` codificado em
        base64).
    """
    if caminho_saida is not None and not caminho_saida.strip():
        raise ValueError("caminho_saida, se informado, não pode ser vazio")
    if caminho_saida is not None and not caminho_saida.lower().endswith(".xlsx"):
        raise ValueError("caminho_saida, se informado, deve terminar em '.xlsx'")

    _validar_modelo(modelo)
    _validar_memoria(memoria)
    _validar_resultado_calculo(resultado_calculo)

    logger.info(
        "gerar_relatorio_excel modelo=%s objeto=%r quantidade_fontes=%d caminho_saida=%r",
        modelo, memoria.get("objeto"), memoria.get("quantidade_fontes"), caminho_saida,
    )

    wb = Workbook()
    _montar_aba_mapa_pesquisa(wb, memoria, resultado_calculo)
    _montar_aba_fontes(wb, memoria)
    _montar_aba_analise_estatistica(wb, resultado_calculo)
    _montar_aba_avisos(wb, memoria, resultado_calculo)

    saida = gravar_documento(
        wb,
        caminho_saida=caminho_saida,
        extensao="xlsx",
        prefixo_nome="relatorio_pesquisa_precos",
    )
    return {
        **saida,
        "modelo": modelo,
        "abas": [ws.title for ws in wb.worksheets],
    }
