"""Tool MCP: gera um documento Word real (.docx) de formalização da pesquisa de preços.

Suporta múltiplos **modelos** de relatório (parâmetro `modelo`, ver
``domain.dominios.MODELOS_RELATORIO_PRECOS``):

- ``"in65"`` (padrão, único implementado): segue os 8 elementos mínimos que
  o **Art. 3º da IN SEGES/ME nº 65/2021** exige que a pesquisa de preços
  contenha, confirmados lendo o texto oficial da norma em
  https://www.gov.br/compras/pt-br/acesso-a-informacao/legislacao/instrucoes-normativas/instrucao-normativa-seges-me-no-65-de-7-de-julho-de-2021 —
  não uma estrutura de "seções 7.1 a 7.9" (essa numeração, usada por
  algumas ferramentas, não existe no texto real da IN, que é corrido em 11
  artigos, sem anexo com seções numeradas dessa forma).
- ``"tic_licitacao"`` e ``"tic_inexigibilidade"``: modelos específicos para
  contratação de itens de TIC (respectivamente via licitação e via
  inexigibilidade/dispensa — ver Arts. 7º e 8º da IN 65/2021), **preparados
  no domínio de valores válidos mas ainda não implementados** — selecioná-
  los levanta ``NotImplementedError`` até que o conteúdo real desses
  modelos seja definido.

Consome o mesmo par de resultados que ``gerar_relatorio_excel`` (o
dicionário devolvido por ``gerar_memoria_pesquisa`` e o dicionário devolvido
por ``calcular_preco_referencia``).

Campos que este servidor não coleta automaticamente (data/hora exata de
acesso a cada portal, cotações de fornecedor fora do que já foi informado em
``fontes``, assinaturas dos responsáveis) são marcados no documento como
``[PENDÊNCIA: ...]`` em vermelho — para revisão humana antes de anexar ao
processo. Este documento é sempre um RASCUNHO estruturado, nunca o documento
final: precisa de revisão e assinatura humana.

A tabela do elemento III (Caracterização das fontes
consultadas) exige uma evidência de acesso por fonte: a coluna "URL /
Evidência de acesso" mostra o ``url`` da fonte (ex.: o link da página do
PNCP/Compras.gov.br do contrato/ata/edital) ou, na ausência dele, o texto
livre de ``evidencia_acesso`` (para fontes sem link, como cotação direta por
e-mail/telefone); se nenhum dos dois estiver preenchido, a própria célula
mostra ``[PENDÊNCIA: ...]`` — nunca fica em branco silenciosamente — e uma
pendência agregada é somada ao final da seção. Pedido explícito do usuário:
todo relatório precisa trazer o link do PNCP/Compras.gov.br do
contrato/ata ou, quando não houver, evidência clara de como acessar a
fonte.

Aceita um parâmetro opcional ``estudo_mercado`` — o
dicionário devolvido por ``gerar_estudo_mercado_tic`` — que insere uma
seção adicional "Estudo de mercado — comparabilidade de soluções técnicas"
entre os elementos II e III, com a matriz de comparação de soluções
candidatas e a justificativa da escolha (Art. 11, VI, da IN SGD/ME nº
94/2022). Essa seção é ADICIONAL aos 8 elementos do Art. 3º da IN 65/2021
— relevante só para contratações de TIC que passaram pela etapa de Estudo
de Mercado do ETP, por isso opcional e fora da numeração I-VIII.
"""
from __future__ import annotations

from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from ..domain.dominios import MODELOS_RELATORIO_PRECOS, MODELOS_RELATORIO_PRECOS_IMPLEMENTADOS
from ..logging_utils import get_logger
from ..utils.arquivo_saida import gravar_documento

logger = get_logger(__name__)

_CAMPOS_OBRIGATORIOS_MEMORIA = (
    "objeto",
    "agente_responsavel",
    "data_pesquisa",
    "metodologia",
    "fontes_consultadas",
    "quantidade_fontes",
)
_CAMPOS_OBRIGATORIOS_RESULTADO_CALCULO = (
    "metodo_utilizado",
    "preco_referencia",
    "estatisticas",
    "amostra_original",
    "amostra_saneada",
    "valores_excluidos",
    "iteracoes_saneamento",
)
_CAMPOS_OBRIGATORIOS_ESTUDO_MERCADO = ("solucoes", "criterios_comparados", "matriz")

_COR_PENDENCIA = RGBColor(0xC0, 0x00, 0x00)

# Heurística para detectar se alguma fonte veio de pesquisa direta com
# fornecedores (Art. 5º, IV, da IN 65/2021) — só nesse caso o Art. 3º, VIII
# exige justificativa da escolha dos fornecedores.
_TERMOS_PESQUISA_DIRETA = ("cotação direta", "cotacao direta", "pesquisa direta")

_NOMES_METODO = {
    "media_saneada": "média, após tratamento/saneamento dos valores (Art. 6º, caput e §3º)",
    "escore_z": "média, após tratamento/saneamento dos valores pelo método Escore Z (Art. 6º, caput e §3º; IPPC-TIC, AGU/MGI)",
    "mediana": "mediana (Art. 6º, caput)",
    "minimo": "menor valor obtido (Art. 6º, caput)",
}


def _validar_memoria(memoria: dict[str, Any]) -> None:
    faltantes = [c for c in _CAMPOS_OBRIGATORIOS_MEMORIA if c not in memoria]
    if faltantes:
        raise ValueError(
            f"memoria não parece vir de gerar_memoria_pesquisa (faltam: {faltantes})"
        )


def _validar_resultado_calculo(resultado_calculo: dict[str, Any]) -> None:
    faltantes = [c for c in _CAMPOS_OBRIGATORIOS_RESULTADO_CALCULO if c not in resultado_calculo]
    if faltantes:
        raise ValueError(
            f"resultado_calculo não parece vir de calcular_preco_referencia (faltam: {faltantes})"
        )


def _validar_estudo_mercado(estudo_mercado: dict[str, Any]) -> None:
    faltantes = [c for c in _CAMPOS_OBRIGATORIOS_ESTUDO_MERCADO if c not in estudo_mercado]
    if faltantes:
        raise ValueError(
            f"estudo_mercado não parece vir de gerar_estudo_mercado_tic (faltam: {faltantes})"
        )


def _validar_modelo(modelo: str) -> None:
    if modelo not in MODELOS_RELATORIO_PRECOS:
        raise ValueError(f"modelo deve ser um de {MODELOS_RELATORIO_PRECOS}")
    if modelo not in MODELOS_RELATORIO_PRECOS_IMPLEMENTADOS:
        raise NotImplementedError(
            f"modelo={modelo!r} está preparado (é um valor válido de "
            "domain.dominios.MODELOS_RELATORIO_PRECOS) mas ainda não foi "
            "implementado — só 'in65' gera relatório hoje. Ver README/CHANGELOG."
        )


_MARCADOR_SEM_EVIDENCIA = "[PENDÊNCIA: informar URL (PNCP/Compras.gov.br) ou evidência de acesso]"


def _url_ou_evidencia(fonte: dict[str, Any]) -> str:
    """URL da fonte (preferencial) ou o texto livre de `evidencia_acesso`;
    se nenhum dos dois estiver preenchido, devolve um marcador de pendência
    visível diretamente na tabela — nunca uma célula em branco (Art. 3º,
    III, da IN SEGES/ME nº 65/2021 exige caracterizar as fontes)."""
    return fonte.get("url") or fonte.get("evidencia_acesso") or _MARCADOR_SEM_EVIDENCIA


def _tem_pesquisa_direta(memoria: dict[str, Any]) -> bool:
    for fonte in memoria.get("fontes_consultadas") or []:
        meio = (fonte.get("meio_consulta") or "").lower()
        if any(termo in meio for termo in _TERMOS_PESQUISA_DIRETA):
            return True
    return False


def _add_pendencia(doc: Document, texto: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(f"[PENDÊNCIA: {texto}]")
    run.bold = True
    run.font.color.rgb = _COR_PENDENCIA


def _add_rotulo_valor(doc: Document, rotulo: str, valor: Any) -> None:
    p = doc.add_paragraph()
    run_rotulo = p.add_run(f"{rotulo}: ")
    run_rotulo.bold = True
    p.add_run("" if valor is None else str(valor))


def _add_tabela(doc: Document, cabecalho: list[str], linhas: list[list[Any]]) -> None:
    tabela = doc.add_table(rows=1, cols=len(cabecalho))
    tabela.style = "Light Grid Accent 1"
    celulas_cabecalho = tabela.rows[0].cells
    for i, titulo in enumerate(cabecalho):
        celulas_cabecalho[i].text = titulo
        for paragrafo in celulas_cabecalho[i].paragraphs:
            for run in paragrafo.runs:
                run.bold = True
    for linha in linhas:
        celulas = tabela.add_row().cells
        for i, valor in enumerate(linha):
            celulas[i].text = "" if valor is None else str(valor)


def _identificacao_processo(doc: Document, memoria: dict[str, Any], numero_processo: str | None) -> None:
    doc.add_heading("Identificação do processo", level=1)
    if numero_processo:
        _add_rotulo_valor(doc, "Número do processo", numero_processo)
    else:
        _add_pendencia(doc, "informar o número do processo administrativo")
    _add_rotulo_valor(doc, "Agente responsável pela pesquisa", memoria.get("agente_responsavel"))
    _add_rotulo_valor(doc, "Data da pesquisa", memoria.get("data_pesquisa"))


def _elemento_1_objeto(doc: Document, memoria: dict[str, Any]) -> None:
    doc.add_heading("I – Descrição do objeto a ser contratado", level=1)
    doc.add_paragraph(str(memoria.get("objeto") or ""))


def _elemento_2_agente(doc: Document, memoria: dict[str, Any]) -> None:
    doc.add_heading("II – Identificação do(s) agente(s) responsável(is) pela pesquisa", level=1)
    doc.add_paragraph(str(memoria.get("agente_responsavel") or ""))


_TITULO_SECAO_ESTUDO_MERCADO = (
    "Estudo de mercado — comparabilidade de soluções técnicas "
    "(TIC — Art. 11, VI, IN SGD/ME nº 94/2022)"
)


def _secao_estudo_mercado(doc: Document, estudo_mercado: dict[str, Any]) -> None:
    doc.add_heading(_TITULO_SECAO_ESTUDO_MERCADO, level=1)
    doc.add_paragraph(
        "Avaliação das alternativas de solução tecnicamente viáveis e, havendo "
        "mais de uma, justificativa técnica e econômica da escolha do tipo de "
        "solução (Art. 11, VI, da IN SGD/ME nº 94/2022)."
    )

    solucoes = estudo_mercado.get("solucoes") or []
    criterios = estudo_mercado.get("criterios_comparados") or []
    matriz = estudo_mercado.get("matriz") or []
    nomes = [solucao.get("nome") for solucao in solucoes]

    if nomes and criterios:
        cabecalho = ["Critério", *nomes]
        linhas = [
            [linha.get("criterio"), *[linha.get(nome) for nome in nomes]]
            for linha in matriz
        ]
        _add_tabela(doc, cabecalho, linhas)
    else:
        doc.add_paragraph("Nenhuma solução/critério informado para comparação.")

    solucao_recomendada = estudo_mercado.get("solucao_recomendada")
    justificativa = estudo_mercado.get("justificativa_escolha")
    if solucao_recomendada:
        _add_rotulo_valor(doc, "Solução recomendada", solucao_recomendada)
        if justificativa:
            doc.add_paragraph(justificativa)
        else:
            _add_pendencia(
                doc,
                "informar a justificativa técnica/econômica da escolha da solução "
                "(Art. 11, VI, da IN SGD/ME nº 94/2022)"
            )
    else:
        _add_pendencia(
            doc,
            "indicar a solução recomendada dentre as avaliadas, com justificativa "
            "(Art. 11, VI, da IN SGD/ME nº 94/2022)"
        )

    for aviso in estudo_mercado.get("avisos") or []:
        doc.add_paragraph(aviso, style="List Bullet")


def _elemento_3_fontes(doc: Document, memoria: dict[str, Any]) -> None:
    doc.add_heading("III – Caracterização das fontes consultadas", level=1)
    doc.add_paragraph(
        "Hierarquia de parâmetros do Art. 5º da IN SEGES/ME nº 65/2021, priorizando "
        "os incisos I e II (§1º): I – Painel de Preços/bancos oficiais (mediana); "
        "II – contratações similares da Administração Pública (até 12 meses); "
        "III – mídia especializada/tabela de referência (até 6 meses); "
        "IV – pesquisa direta com no mínimo 3 fornecedores (até 6 meses); "
        "V – base nacional de notas fiscais eletrônicas (até 12 meses)."
    )
    fontes = memoria.get("fontes_consultadas") or []
    if fontes:
        cabecalho = [
            "Informante", "CNPJ", "Meio de consulta", "Valor unitário (R$)",
            "Quantidade", "Data de referência", "Especificação", "Condições",
            "URL / Evidência de acesso",
        ]
        linhas = [
            [
                fonte.get("informante"),
                fonte.get("cnpj"),
                fonte.get("meio_consulta"),
                fonte.get("valor_unitario"),
                fonte.get("quantidade"),
                fonte.get("data_referencia"),
                fonte.get("especificacao"),
                fonte.get("condicoes"),
                _url_ou_evidencia(fonte),
            ]
            for fonte in fontes
        ]
        _add_tabela(doc, cabecalho, linhas)
        sem_evidencia = sum(
            1 for fonte in fontes if not fonte.get("url") and not fonte.get("evidencia_acesso")
        )
        if sem_evidencia:
            _add_pendencia(
                doc,
                f"{sem_evidencia} de {len(fontes)} fonte(s) sem link (PNCP/Compras.gov.br) "
                "nem evidência de acesso registrada — ver coluna 'URL / Evidência de acesso' "
                "na tabela acima. Exigido pelo Art. 3º, III, da IN SEGES/ME nº 65/2021."
            )
    else:
        doc.add_paragraph("Nenhuma fonte informada.")
    _add_pendencia(doc, "confirmar data e hora exatas de acesso a cada portal consultado")

    buscas_frustradas = memoria.get("buscas_frustradas") or []
    if buscas_frustradas:
        doc.add_heading("Buscas sem resultado", level=2)
        for busca in buscas_frustradas:
            termo = f' — termo: "{busca.get("termo_busca")}"' if busca.get("termo_busca") else ""
            doc.add_paragraph(
                f"{busca.get('meio_consulta')}{termo} — motivo: {busca.get('motivo')}",
                style="List Bullet",
            )


def _elemento_4_serie_precos(doc: Document, resultado_calculo: dict[str, Any]) -> None:
    doc.add_heading("IV – Série de preços coletados", level=1)
    amostra_original = (resultado_calculo.get("amostra_original") or {}).get("valores") or []
    valores_excluidos = list(resultado_calculo.get("valores_excluidos") or [])
    cabecalho = ["#", "Valor (R$)", "Situação"]
    linhas = []
    excluidos_restantes = list(valores_excluidos)
    for i, valor in enumerate(amostra_original, start=1):
        situacao = "Considerado no cálculo"
        if valor in excluidos_restantes:
            situacao = "Desconsiderado (Art. 6º, caput/§3º)"
            excluidos_restantes.remove(valor)
        linhas.append([i, valor, situacao])
    _add_tabela(doc, cabecalho, linhas)


def _elemento_5_metodo(doc: Document, resultado_calculo: dict[str, Any]) -> None:
    doc.add_heading("V – Método estatístico aplicado para a definição do valor estimado", level=1)
    metodo = resultado_calculo.get("metodo_utilizado")
    nome_metodo = _NOMES_METODO.get(metodo, str(metodo))
    doc.add_paragraph(f"Método aplicado: {nome_metodo}.")

    estatisticas = resultado_calculo.get("estatisticas") or {}
    _add_tabela(
        doc,
        ["Média (R$)", "Mediana (R$)", "Mínimo (R$)", "Máximo (R$)", "Desvio-padrão", "CV (%)"],
        [[
            estatisticas.get("media"),
            estatisticas.get("mediana"),
            estatisticas.get("minimo"),
            estatisticas.get("maximo"),
            estatisticas.get("desvio_padrao"),
            estatisticas.get("coeficiente_variacao_percentual"),
        ]],
    )


def _elemento_6_justificativa_metodologia(
    doc: Document,
    memoria: dict[str, Any],
    resultado_calculo: dict[str, Any],
    justificativa_metodologia: str | None,
) -> None:
    doc.add_heading("VI – Justificativas para a metodologia utilizada", level=1)
    doc.add_paragraph(
        "Art. 6º, §3º, da IN SEGES/ME nº 65/2021: para desconsideração dos "
        "valores inexequíveis, inconsistentes ou excessivamente elevados, "
        "devem ser adotados critérios fundamentados e descritos no processo."
    )
    valores_excluidos = resultado_calculo.get("valores_excluidos") or []
    iteracoes = resultado_calculo.get("iteracoes_saneamento")
    if valores_excluidos:
        _add_rotulo_valor(
            doc,
            "Valores desconsiderados automaticamente (fora de média ± desvio-padrão)",
            ", ".join(str(v) for v in valores_excluidos),
        )
        _add_rotulo_valor(doc, "Iterações de saneamento aplicadas", iteracoes)
    else:
        doc.add_paragraph("Nenhum valor foi desconsiderado da amostra coletada.")

    avisos = list(memoria.get("avisos") or []) + list(resultado_calculo.get("avisos") or [])
    for aviso in avisos:
        doc.add_paragraph(aviso, style="List Bullet")

    if justificativa_metodologia:
        doc.add_paragraph(justificativa_metodologia)
    else:
        _add_pendencia(
            doc,
            "acrescentar, se necessário, justificativa adicional para a metodologia "
            "aplicada além da desconsideração automática de outliers"
        )


def _elemento_7_memoria_calculo(
    doc: Document, resultado_calculo: dict[str, Any], quantidade_estimada: float | None
) -> None:
    doc.add_heading("VII – Memória de cálculo do valor estimado e documentos que lhe dão suporte", level=1)
    preco = resultado_calculo.get("preco_referencia")
    p = doc.add_paragraph()
    p.add_run("Valor estimado (preço de referência unitário): ").bold = True
    run_valor = p.add_run(f"R$ {preco:.4f}" if isinstance(preco, (int, float)) else str(preco))
    run_valor.bold = True
    run_valor.font.size = Pt(12)

    if quantidade_estimada is not None and isinstance(preco, (int, float)):
        valor_total = preco * quantidade_estimada
        _add_rotulo_valor(doc, "Quantidade estimada para a contratação", quantidade_estimada)
        _add_rotulo_valor(doc, "Valor total estimado (R$)", f"{valor_total:.2f}")
    else:
        _add_pendencia(doc, "informar a quantidade estimada da contratação para calcular o valor total")

    _add_pendencia(doc, "anexar os documentos de suporte da memória de cálculo (capturas de tela, propostas, notas fiscais etc.)")


def _elemento_8_justificativa_fornecedores(
    doc: Document, memoria: dict[str, Any], justificativa_fornecedores: str | None
) -> None:
    doc.add_heading("VIII – Justificativa da escolha dos fornecedores", level=1)
    doc.add_paragraph(
        "Exigida pelo Art. 3º, VIII, apenas quando a pesquisa usa pesquisa "
        "direta com fornecedores (Art. 5º, IV)."
    )
    if not _tem_pesquisa_direta(memoria):
        doc.add_paragraph(
            "Não aplicável — nenhuma fonte de pesquisa direta com fornecedores "
            "(Art. 5º, IV) foi utilizada nesta pesquisa."
        )
        return
    if justificativa_fornecedores:
        doc.add_paragraph(justificativa_fornecedores)
    else:
        _add_pendencia(
            doc,
            "informar a justificativa da escolha dos fornecedores consultados por "
            "pesquisa direta (Art. 5º, IV e §2º — prazo de resposta, propostas "
            "formais com CPF/CNPJ, endereço, data de emissão e responsável)"
        )


def _responsaveis(doc: Document, responsaveis: list[str] | None) -> None:
    doc.add_heading("Responsáveis pela pesquisa", level=1)
    nomes = responsaveis or []
    cabecalho = ["Nome / matrícula / cargo", "Assinatura", "Data"]
    if nomes:
        linhas = [[nome, "", ""] for nome in nomes]
    else:
        linhas = [["", "", ""]]
    _add_tabela(doc, cabecalho, linhas)
    if not nomes:
        _add_pendencia(doc, "informar nome, matrícula e cargo dos responsáveis pela pesquisa")
    _add_pendencia(doc, "coletar assinatura e data dos responsáveis antes de anexar ao processo")


async def gerar_relatorio_word(
    memoria: dict[str, Any],
    resultado_calculo: dict[str, Any],
    caminho_saida: str | None = None,
    modelo: str = "in65",
    numero_processo: str | None = None,
    responsaveis: list[str] | None = None,
    quantidade_estimada: float | None = None,
    justificativa_metodologia: str | None = None,
    justificativa_fornecedores: str | None = None,
    estudo_mercado: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Gera um documento Word (.docx) real de formalização da pesquisa de preços.

    Complementa ``gerar_memoria_pesquisa`` e ``gerar_relatorio_excel``: esta
    tool grava um arquivo ``.docx`` de verdade. Para ``modelo="in65"``
    (padrão, único implementado), a estrutura segue os 8 elementos que o
    **Art. 3º da IN SEGES/ME nº 65/2021** exige que a formalização da
    pesquisa de preços contenha, no mínimo: (I) descrição do objeto,
    (II) identificação do agente responsável, (III) caracterização das
    fontes consultadas, (IV) série de preços coletados, (V) método
    estatístico aplicado, (VI) justificativas para a metodologia utilizada,
    (VII) memória de cálculo do valor estimado e documentos de suporte, e
    (VIII) justificativa da escolha dos fornecedores — só quando a pesquisa
    usa pesquisa direta com fornecedores (Art. 5º, IV). Note que essa NÃO é
    uma estrutura de "seções 7.1 a 7.9": essa numeração, usada por algumas
    ferramentas, não existe no texto oficial da norma, que não tem anexo
    com seções numeradas dessa forma.

    ``modelo="tic_licitacao"`` e ``modelo="tic_inexigibilidade"`` (modelos
    específicos para contratação de TIC via licitação e via inexigibilidade/
    dispensa, respectivamente) estão **preparados como valores válidos mas
    ainda não implementados** — selecioná-los levanta ``NotImplementedError``.

    Campos que este servidor não coleta automaticamente (data/hora exata de
    acesso a cada portal, comprovantes, assinaturas) aparecem destacados em
    vermelho como ``[PENDÊNCIA: ...]`` — o documento gerado é sempre um
    RASCUNHO estruturado, que precisa de revisão e assinatura humana antes
    de ser anexado ao processo administrativo.

    Esta tool é **quase pura**: não consulta nenhuma API — sua única
    operação com efeito colateral é gravar/serializar o documento (ver
    ``caminho_saida`` abaixo).

    Args:
        memoria: O dicionário devolvido por ``gerar_memoria_pesquisa``.
        resultado_calculo: O dicionário devolvido por
            ``calcular_preco_referencia`` para a MESMA amostra usada na
            memória (mesmo requisito de ``gerar_relatorio_excel``).
        caminho_saida: Caminho absoluto ou relativo do arquivo ``.docx`` a
            gravar no disco de onde o servidor roda (diretórios
            intermediários são criados automaticamente). Use este
            parâmetro quando o servidor roda em ``stdio`` (mesma máquina do
            usuário, ex.: Claude Desktop) ou via túnel a partir da própria
            máquina (ex.: ChatGPT) — nesses casos o disco do servidor É o
            disco do usuário. **Omita este parâmetro quando o servidor
            estiver hospedado remotamente** (ex.: custom connector
            compartilhado por vários usuários): o documento é então
            serializado inteiramente em memória, sem tocar o disco do
            servidor, e devolvido em base64 (ver Returns) com um nome de
            arquivo único — cabe ao cliente decodificar e salvar do lado
            do usuário.
        modelo: Modelo de relatório — "in65" (padrão, implementado),
            "tic_licitacao" ou "tic_inexigibilidade" (preparados, ainda não
            implementados — levantam ``NotImplementedError``).
        numero_processo: Número do processo administrativo. Se omitido, a
            identificação do processo é marcada como pendência.
        responsaveis: Nomes (idealmente com matrícula/cargo) para a tabela
            de responsáveis. Se omitido, a tabela sai com uma linha em
            branco e uma pendência.
        quantidade_estimada: Quantidade estimada da contratação, para
            calcular o valor total estimado no elemento VII (preço de
            referência × quantidade). Se omitido, marcado como pendência.
        justificativa_metodologia: Texto livre com justificativa adicional
            para a metodologia (Art. 6º, §3º), além da desconsideração
            automática de outliers já registrada no elemento VI.
        justificativa_fornecedores: Texto livre com a justificativa da
            escolha dos fornecedores (Art. 3º, VIII e Art. 5º, IV) — só
            relevante quando alguma fonte usa pesquisa direta; ignorado
            (com nota de "não aplicável") caso contrário.
        estudo_mercado: O dicionário devolvido por
            ``gerar_estudo_mercado_tic``, opcional. Quando informado,
            insere uma seção adicional "Estudo de mercado —
            comparabilidade de soluções técnicas" entre os elementos II e
            III, com a matriz de comparação de soluções candidatas e a
            justificativa da escolha (Art. 11, VI, da IN SGD/ME nº
            94/2022) — relevante só para contratações de TIC que
            passaram pela etapa de Estudo de Mercado do ETP.

    Returns:
        Dicionário com ``modelo``, ``elementos`` (títulos dos elementos I a
        VIII do Art. 3º, na ordem em que aparecem), ``secoes_adicionais``
        (títulos de seções extras incluídas além dos elementos I-VIII —
        hoje só a de estudo de mercado, quando ``estudo_mercado`` é
        informado), ``tamanho_bytes`` e mais: se ``caminho_saida`` foi
        informado, ``caminho_saida`` (caminho absoluto do arquivo
        gravado); se foi omitido, ``nome_arquivo`` (nome sugerido, com um
        UUID) e ``conteudo_base64`` (conteúdo do ``.docx`` codificado em
        base64).
    """
    if caminho_saida is not None and not caminho_saida.strip():
        raise ValueError("caminho_saida, se informado, não pode ser vazio")
    if caminho_saida is not None and not caminho_saida.lower().endswith(".docx"):
        raise ValueError("caminho_saida, se informado, deve terminar em '.docx'")

    _validar_modelo(modelo)
    _validar_memoria(memoria)
    _validar_resultado_calculo(resultado_calculo)
    if estudo_mercado is not None:
        _validar_estudo_mercado(estudo_mercado)

    logger.info(
        "gerar_relatorio_word modelo=%s objeto=%r quantidade_fontes=%d caminho_saida=%r",
        modelo, memoria.get("objeto"), memoria.get("quantidade_fontes"), caminho_saida,
    )

    doc = Document()
    titulo = doc.add_heading(f"Pesquisa de Preços — {memoria.get('objeto')}", level=0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "Formalizado conforme os elementos mínimos do Art. 3º da "
        "IN SEGES/ME nº 65/2021."
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER

    elementos = [
        "I – Descrição do objeto a ser contratado",
        "II – Identificação do(s) agente(s) responsável(is) pela pesquisa",
        "III – Caracterização das fontes consultadas",
        "IV – Série de preços coletados",
        "V – Método estatístico aplicado para a definição do valor estimado",
        "VI – Justificativas para a metodologia utilizada",
        "VII – Memória de cálculo do valor estimado e documentos que lhe dão suporte",
        "VIII – Justificativa da escolha dos fornecedores",
    ]

    secoes_adicionais: list[str] = []

    _identificacao_processo(doc, memoria, numero_processo)
    _elemento_1_objeto(doc, memoria)
    _elemento_2_agente(doc, memoria)
    if estudo_mercado is not None:
        _secao_estudo_mercado(doc, estudo_mercado)
        secoes_adicionais.append(_TITULO_SECAO_ESTUDO_MERCADO)
    _elemento_3_fontes(doc, memoria)
    _elemento_4_serie_precos(doc, resultado_calculo)
    _elemento_5_metodo(doc, resultado_calculo)
    _elemento_6_justificativa_metodologia(doc, memoria, resultado_calculo, justificativa_metodologia)
    _elemento_7_memoria_calculo(doc, resultado_calculo, quantidade_estimada)
    _elemento_8_justificativa_fornecedores(doc, memoria, justificativa_fornecedores)
    _responsaveis(doc, responsaveis)

    saida = gravar_documento(
        doc,
        caminho_saida=caminho_saida,
        extensao="docx",
        prefixo_nome="relatorio_pesquisa_precos",
    )
    return {
        **saida,
        "modelo": modelo,
        "elementos": elementos,
        "secoes_adicionais": secoes_adicionais,
    }
