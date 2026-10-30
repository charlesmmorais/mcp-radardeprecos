"""Testes de `gerar_relatorio_word` (tool quase pura — só grava um arquivo).

A estrutura testada segue os 8 elementos mínimos do Art. 3º da IN SEGES/ME
nº 65/2021 (I a VIII) — não a estrutura de "seções 7.1 a 7.9" usada até a
v0.1.22, que era herdada de um template de terceiros e não existe no texto
oficial da norma.
"""
from __future__ import annotations

import base64
import io

import pytest
from docx import Document

from mcp_radardeprecos.tools.gerar_relatorio_word import gerar_relatorio_word

_RESULTADO_CALCULO = {
    "metodo_utilizado": "media_saneada",
    "preco_referencia": 98.4,
    "amostra_original": {"quantidade": 6, "valores": [100.0, 95.0, 98.0, 102.0, 97.0, 350.0]},
    "amostra_saneada": [100.0, 95.0, 98.0, 102.0, 97.0],
    "valores_excluidos": [350.0],
    "iteracoes_saneamento": 1,
    "estatisticas": {
        "media": 98.4,
        "mediana": 98.0,
        "minimo": 95.0,
        "maximo": 102.0,
        "desvio_padrao": 2.42,
        "coeficiente_variacao_percentual": 2.46,
    },
    "avisos": [],
}

_MEMORIA = {
    "objeto": "Notebook corporativo",
    "agente_responsavel": "Fulano de Tal - Analista",
    "data_pesquisa": "2026-07-04",
    "metodologia": {
        "metodo_utilizado": "media_saneada",
        "preco_referencia": 98.4,
        "estatisticas": _RESULTADO_CALCULO["estatisticas"],
    },
    "fontes_consultadas": [
        {
            "informante": "Fornecedor A",
            "cnpj": "11111111000111",
            "meio_consulta": "PNCP - contratação",
            "valor_unitario": 100.0,
            "quantidade": 5,
            "data_referencia": "2026-05-01",
            "especificacao": "Notebook i7 16GB",
            "condicoes": "Entrega em 30 dias",
            "url": "https://pncp.gov.br/exemplo",
        },
        {
            "informante": "Fornecedor B",
            "meio_consulta": "Compras.gov.br",
            "valor_unitario": 95.0,
            "data_referencia": "2026-05-02",
        },
    ],
    "quantidade_fontes": 2,
    "buscas_frustradas": [
        {"meio_consulta": "Painel de Preços", "termo_busca": "notebook", "motivo": "sem resultado"}
    ],
    "avisos": ["Amostra pequena — considerar ampliar."],
}

_MEMORIA_COM_PESQUISA_DIRETA = {
    **_MEMORIA,
    "fontes_consultadas": [
        *_MEMORIA["fontes_consultadas"],
        {
            "informante": "Fornecedor C",
            "meio_consulta": "Cotação direta",
            "valor_unitario": 99.0,
            "data_referencia": "2026-05-03",
        },
    ],
}

_ELEMENTOS_ESPERADOS = [
    "I – Descrição do objeto a ser contratado",
    "II – Identificação do(s) agente(s) responsável(is) pela pesquisa",
    "III – Caracterização das fontes consultadas",
    "IV – Série de preços coletados",
    "V – Método estatístico aplicado para a definição do valor estimado",
    "VI – Justificativas para a metodologia utilizada",
    "VII – Memória de cálculo do valor estimado e documentos que lhe dão suporte",
    "VIII – Justificativa da escolha dos fornecedores",
]


def _texto_completo(doc: Document) -> str:
    partes = [p.text for p in doc.paragraphs]
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                partes.append(celula.text)
    return "\n".join(partes)


async def test_rejeita_caminho_sem_extensao_docx(tmp_path):
    with pytest.raises(ValueError, match="docx"):
        await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO, str(tmp_path / "saida.txt"))


async def test_rejeita_caminho_vazio():
    with pytest.raises(ValueError):
        await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO, "")


async def test_rejeita_memoria_sem_campos_esperados(tmp_path):
    with pytest.raises(ValueError, match="gerar_memoria_pesquisa"):
        await gerar_relatorio_word({"algo": 1}, _RESULTADO_CALCULO, str(tmp_path / "saida.docx"))


async def test_rejeita_resultado_calculo_sem_campos_esperados(tmp_path):
    with pytest.raises(ValueError, match="calcular_preco_referencia"):
        await gerar_relatorio_word(_MEMORIA, {"algo": 1}, str(tmp_path / "saida.docx"))


async def test_rejeita_modelo_invalido(tmp_path):
    with pytest.raises(ValueError, match="modelo"):
        await gerar_relatorio_word(
            _MEMORIA, _RESULTADO_CALCULO, str(tmp_path / "saida.docx"), modelo="xyz"
        )


@pytest.mark.parametrize("modelo", ["tic_licitacao", "tic_inexigibilidade"])
async def test_modelos_tic_preparados_mas_nao_implementados(tmp_path, modelo):
    with pytest.raises(NotImplementedError, match=modelo):
        await gerar_relatorio_word(
            _MEMORIA, _RESULTADO_CALCULO, str(tmp_path / "saida.docx"), modelo=modelo
        )


async def test_gera_arquivo_com_elementos_na_ordem_esperada(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    resultado = await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO, str(caminho))

    assert caminho.exists()
    assert resultado["modelo"] == "in65"
    assert resultado["elementos"] == _ELEMENTOS_ESPERADOS
    assert resultado["tamanho_bytes"] == caminho.stat().st_size


async def test_cria_diretorios_intermediarios(tmp_path):
    caminho = tmp_path / "sub1" / "sub2" / "relatorio.docx"
    resultado = await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO, str(caminho))
    assert caminho.exists()
    assert resultado["caminho_saida"] == str(caminho)


async def test_identificacao_processo_marca_pendencia_quando_sem_numero_processo(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO, str(caminho))
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "PENDÊNCIA" in texto
    assert "número do processo" in texto


async def test_identificacao_processo_usa_numero_processo_quando_informado(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    await gerar_relatorio_word(
        _MEMORIA, _RESULTADO_CALCULO, str(caminho), numero_processo="23000.000123/2026-01"
    )
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "23000.000123/2026-01" in texto


async def test_elemento_3_lista_fontes_hierarquia_e_buscas_frustradas(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO, str(caminho))
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "Fornecedor A" in texto
    assert "Fornecedor B" in texto
    assert "https://pncp.gov.br/exemplo" in texto
    assert "Art. 5º" in texto
    assert "Buscas sem resultado" in texto


async def test_elemento_3_marca_pendencia_quando_fonte_sem_url_e_sem_evidencia(tmp_path):
    # "Fornecedor B" em _MEMORIA não tem "url" nem "evidencia_acesso".
    caminho = tmp_path / "relatorio.docx"
    await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO, str(caminho))
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "URL / Evidência de acesso" in texto
    assert "[PENDÊNCIA: informar URL (PNCP/Compras.gov.br) ou evidência de acesso]" in texto
    assert "1 de 2 fonte(s) sem link" in texto
    assert "Art. 3º, III" in texto


async def test_elemento_3_evidencia_acesso_aparece_quando_sem_url(tmp_path):
    memoria = {
        **_MEMORIA,
        "fontes_consultadas": [
            {
                "informante": "Fornecedor Direto",
                "meio_consulta": "Cotação direta",
                "valor_unitario": 97.0,
                "data_referencia": "2026-05-04",
                "evidencia_acesso": "E-mail de cotação anexado ao processo, fls. 20-22",
            }
        ],
        "quantidade_fontes": 1,
    }
    caminho = tmp_path / "relatorio.docx"
    await gerar_relatorio_word(memoria, _RESULTADO_CALCULO, str(caminho))
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "E-mail de cotação anexado ao processo, fls. 20-22" in texto
    assert "sem link" not in texto  # nenhuma pendência agregada de evidência


async def test_elemento_4_marca_valor_excluido_na_tabela(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO, str(caminho))
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "Desconsiderado (Art. 6º, caput/§3º)" in texto
    assert "350.0" in texto


async def test_elemento_5_nomeia_metodo_aplicado(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO, str(caminho))
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "Art. 6º" in texto


async def test_elemento_6_sem_justificativa_extra_gera_pendencia(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO, str(caminho))
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "350.0" in texto  # valor desconsiderado listado
    assert "justificativa adicional" in texto


async def test_elemento_6_com_justificativa_extra_nao_gera_pendencia_para_ela(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    await gerar_relatorio_word(
        _MEMORIA, _RESULTADO_CALCULO, str(caminho),
        justificativa_metodologia="Valor de R$ 350,00 desconsiderado por incompatibilidade de especificação.",
    )
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "incompatibilidade de especificação" in texto


async def test_elemento_7_sem_quantidade_estimada_gera_pendencia(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO, str(caminho))
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "quantidade estimada" in texto


async def test_elemento_7_com_quantidade_estimada_calcula_valor_total(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    await gerar_relatorio_word(
        _MEMORIA, _RESULTADO_CALCULO, str(caminho), quantidade_estimada=50
    )
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "4920.00" in texto  # 98.4 * 50


async def test_elemento_8_nao_aplicavel_sem_pesquisa_direta(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO, str(caminho))
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "Não aplicável" in texto


async def test_elemento_8_pendencia_quando_ha_pesquisa_direta_sem_justificativa(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    await gerar_relatorio_word(_MEMORIA_COM_PESQUISA_DIRETA, _RESULTADO_CALCULO, str(caminho))
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "justificativa da escolha dos fornecedores" in texto


async def test_elemento_8_usa_justificativa_informada(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    await gerar_relatorio_word(
        _MEMORIA_COM_PESQUISA_DIRETA, _RESULTADO_CALCULO, str(caminho),
        justificativa_fornecedores="Fornecedores escolhidos por atuarem no mesmo ramo do objeto.",
    )
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "atuarem no mesmo ramo do objeto" in texto


async def test_responsaveis_sem_lista_gera_pendencia(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO, str(caminho))
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "nome, matrícula e cargo" in texto


async def test_responsaveis_com_lista_preenche_tabela(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    await gerar_relatorio_word(
        _MEMORIA, _RESULTADO_CALCULO, str(caminho),
        responsaveis=["Fulano de Tal - Analista de Compras"],
    )
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "Fulano de Tal - Analista de Compras" in texto


_ESTUDO_MERCADO = {
    "solucoes": [
        {"nome": "SaaS Fornecedor X", "criterios": {"modelo_licenciamento": "SaaS anual"}},
        {"nome": "On-premise Fornecedor Y", "criterios": {"modelo_licenciamento": "Perpétua"}},
    ],
    "criterios_comparados": ["modelo_licenciamento"],
    "matriz": [
        {
            "criterio": "modelo_licenciamento",
            "SaaS Fornecedor X": "SaaS anual",
            "On-premise Fornecedor Y": "Perpétua",
        },
    ],
    "solucao_recomendada": "SaaS Fornecedor X",
    "justificativa_escolha": "Menor TCO em 36 meses.",
    "avisos": [],
}


async def test_rejeita_estudo_mercado_sem_campos_esperados(tmp_path):
    with pytest.raises(ValueError, match="gerar_estudo_mercado_tic"):
        await gerar_relatorio_word(
            _MEMORIA, _RESULTADO_CALCULO, str(tmp_path / "saida.docx"),
            estudo_mercado={"algo": 1},
        )


async def test_sem_estudo_mercado_nao_inclui_secao(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    resultado = await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO, str(caminho))
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "Estudo de mercado" not in texto
    assert resultado["secoes_adicionais"] == []


async def test_com_estudo_mercado_inclui_secao_e_matriz(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    resultado = await gerar_relatorio_word(
        _MEMORIA, _RESULTADO_CALCULO, str(caminho), estudo_mercado=_ESTUDO_MERCADO,
    )
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "Estudo de mercado" in texto
    assert "Art. 11" in texto
    assert "SaaS Fornecedor X" in texto
    assert "On-premise Fornecedor Y" in texto
    assert "SaaS anual" in texto
    assert "Menor TCO em 36 meses." in texto
    assert resultado["secoes_adicionais"] == [
        "Estudo de mercado — comparabilidade de soluções técnicas "
        "(TIC — Art. 11, VI, IN SGD/ME nº 94/2022)"
    ]


async def test_estudo_mercado_sem_solucao_recomendada_gera_pendencia(tmp_path):
    caminho = tmp_path / "relatorio.docx"
    estudo_sem_recomendacao = {**_ESTUDO_MERCADO, "solucao_recomendada": None, "justificativa_escolha": None}
    await gerar_relatorio_word(
        _MEMORIA, _RESULTADO_CALCULO, str(caminho), estudo_mercado=estudo_sem_recomendacao,
    )
    doc = Document(str(caminho))
    texto = _texto_completo(doc)
    assert "indicar a solução recomendada" in texto


# --- Modo hospedado: caminho_saida omitido -> base64 em memória ---


async def test_sem_caminho_saida_devolve_base64_em_vez_de_gravar_no_disco(tmp_path, monkeypatch):
    # Muda o diretório de trabalho para tmp_path e confirma que nada é
    # gravado nele — o modo sem caminho_saida nunca deve tocar o disco.
    monkeypatch.chdir(tmp_path)

    resultado = await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO)

    assert "caminho_saida" not in resultado
    assert "nome_arquivo" in resultado
    assert resultado["nome_arquivo"].startswith("relatorio_pesquisa_precos_")
    assert resultado["nome_arquivo"].endswith(".docx")
    assert resultado["tamanho_bytes"] > 0
    assert list(tmp_path.iterdir()) == []


async def test_sem_caminho_saida_base64_decodifica_para_docx_valido():
    resultado = await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO)

    conteudo = base64.b64decode(resultado["conteudo_base64"])
    assert len(conteudo) == resultado["tamanho_bytes"]

    doc = Document(io.BytesIO(conteudo))
    texto = _texto_completo(doc)
    assert "Notebook corporativo" in texto


async def test_dois_arquivos_gerados_sem_caminho_saida_tem_nomes_diferentes():
    r1 = await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO)
    r2 = await gerar_relatorio_word(_MEMORIA, _RESULTADO_CALCULO)

    assert r1["nome_arquivo"] != r2["nome_arquivo"]
