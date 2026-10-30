"""Testes de `exportar_resultados` (CSV/XLSX a partir de uma lista de registros)."""
from __future__ import annotations

import base64
import csv
import io

import pytest
from openpyxl import load_workbook

from mcp_radardeprecos.tools.exportar_resultados import exportar_resultados

_REGISTROS = [
    {
        "numero_ata_registro_preco": "001/2026",
        "valor_total": 1000.0,
        "fornecedor": {"cnpj": "12345678000199", "nome": "Fornecedor Teste"},
    },
    {
        "numero_ata_registro_preco": "002/2026",
        "valor_total": 2000.0,
        "fornecedor": {"cnpj": "98765432000188", "nome": "Outro Fornecedor"},
    },
]


async def test_rejeita_formato_invalido():
    with pytest.raises(ValueError, match="formato"):
        await exportar_resultados(_REGISTROS, formato="pdf")


async def test_rejeita_registros_vazio():
    with pytest.raises(ValueError, match="vazio"):
        await exportar_resultados([])


async def test_csv_achata_dicts_aninhados_com_ponto(tmp_path):
    caminho = tmp_path / "saida.csv"
    resultado = await exportar_resultados(_REGISTROS, formato="csv", caminho_saida=str(caminho))
    assert resultado["formato"] == "csv"
    assert resultado["total_registros_exportados"] == 2
    assert "fornecedor.cnpj" in resultado["colunas"]
    assert "fornecedor.nome" in resultado["colunas"]

    conteudo = caminho.read_bytes().decode("utf-8-sig")
    linhas = list(csv.DictReader(io.StringIO(conteudo)))
    assert len(linhas) == 2
    assert linhas[0]["fornecedor.cnpj"] == "12345678000199"
    assert linhas[1]["numero_ata_registro_preco"] == "002/2026"


async def test_xlsx_gera_planilha_valida_com_cabecalho(tmp_path):
    caminho = tmp_path / "saida.xlsx"
    resultado = await exportar_resultados(_REGISTROS, formato="xlsx", caminho_saida=str(caminho))
    assert resultado["formato"] == "xlsx"

    pasta = load_workbook(caminho)
    aba = pasta.active
    cabecalho = [c.value for c in next(aba.iter_rows(min_row=1, max_row=1))]
    assert "fornecedor.cnpj" in cabecalho
    linhas = list(aba.iter_rows(min_row=2, values_only=True))
    assert len(linhas) == 2


async def test_listas_viram_texto_unido_por_ponto_e_virgula(tmp_path):
    registros = [{"id": "1", "tags": ["a", "b", "c"]}]
    caminho = tmp_path / "saida.csv"
    await exportar_resultados(registros, formato="csv", caminho_saida=str(caminho))
    conteudo = caminho.read_bytes().decode("utf-8-sig")
    linhas = list(csv.DictReader(io.StringIO(conteudo)))
    assert linhas[0]["tags"] == "a; b; c"


async def test_nome_colunas_define_ordem_e_subconjunto(tmp_path):
    caminho = tmp_path / "saida.csv"
    resultado = await exportar_resultados(
        _REGISTROS,
        formato="csv",
        caminho_saida=str(caminho),
        nome_colunas=["numero_ata_registro_preco", "valor_total"],
    )
    assert resultado["colunas"] == ["numero_ata_registro_preco", "valor_total"]
    conteudo = caminho.read_bytes().decode("utf-8-sig")
    linhas = list(csv.DictReader(io.StringIO(conteudo)))
    assert set(linhas[0].keys()) == {"numero_ata_registro_preco", "valor_total"}


async def test_itens_nao_dict_sao_ignorados_com_aviso(tmp_path):
    registros = [_REGISTROS[0], "nao é um dict", 42]
    caminho = tmp_path / "saida.csv"
    resultado = await exportar_resultados(registros, formato="csv", caminho_saida=str(caminho))
    assert resultado["total_registros_exportados"] == 1
    assert any("ignorado" in a for a in resultado["avisos"])


async def test_sem_caminho_saida_devolve_base64():
    resultado = await exportar_resultados(_REGISTROS, formato="xlsx")
    assert "caminho_saida" not in resultado
    assert resultado["nome_arquivo"].endswith(".xlsx")
    conteudo = base64.b64decode(resultado["conteudo_base64"])
    pasta = load_workbook(io.BytesIO(conteudo))
    assert pasta.active.max_row == 3  # cabeçalho + 2 registros


async def test_registros_so_com_itens_invalidos_sem_colunas_levanta_erro():
    with pytest.raises(ValueError, match="coluna"):
        await exportar_resultados(["nao é dict", 42])
